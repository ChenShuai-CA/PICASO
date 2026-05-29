import sys
import os
import docx

# Ensure UTF-8 output on Windows
sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"C:\Users\chens\Documents\场景生成研究\stage2"
output_path = r"C:\Users\chens\Documents\场景生成研究\stage2_extracted_all.txt"

files = [
    "C-NCAPE-NCAPISO标准对自动驾驶场景生成的约束条件.docx",
    "Mamba 与选择性状态空间模型在自动驾驶轨迹建模中的最新进展.docx",
    "因果推断与反事实推理在自动驾驶安全验证中的前沿应用.docx",
    "域自适应与域泛化方法在自动驾驶数据迁移中的应用.docx",
    "安全关键场景生成领域 SOTA 方法的定量结果对标.docx",
    "物理信息约束在自动驾驶轨迹场景生成中的最新方法.docx",
    "自动驾驶安全关键场景生成方法的完整技术全景.docx",
    "自动驾驶安全关键场景生成的评估指标体系.docx",
]

with open(output_path, "w", encoding="utf-8") as out:
    for fname in files:
        fpath = os.path.join(base_dir, fname)
        out.write(f"===== {fname} =====\n")

        if not os.path.exists(fpath):
            out.write(f"[ERROR] File not found: {fpath}\n\n")
            continue

        try:
            doc = docx.Document(fpath)

            # Extract paragraphs
            out.write("\n--- Paragraphs ---\n")
            for para in doc.paragraphs:
                out.write(para.text + "\n")

            # Extract tables
            if doc.tables:
                out.write(f"\n--- Tables ({len(doc.tables)} total) ---\n")
                for ti, table in enumerate(doc.tables):
                    out.write(f"\n[Table {ti+1}]\n")
                    for ri, row in enumerate(table.rows):
                        cells = [cell.text.replace("\n", " ") for cell in row.cells]
                        out.write(" | ".join(cells) + "\n")
            out.write("\n")
        except Exception as e:
            out.write(f"[ERROR] Failed to extract: {e}\n\n")

print("Done. Output written to:", output_path)
